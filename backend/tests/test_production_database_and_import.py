import io
import pytest
import openpyxl
from fastapi.testclient import TestClient
from app.main import app
from app.db.session import get_db
from app.db.models.user import User, UserRole
from app.db.models.department import Department
from app.config.security import hash_password

def test_database_status_and_roster_import(db_session):
    admin = User(
        email="admin@printplatform.local",
        password_hash=hash_password("AdminPass123!"),
        role=UserRole.ADMIN,
        is_active=True
    )
    dept = Department(code="CSE", name="Computer Science", description="CS")
    db_session.add(admin)
    db_session.add(dept)
    db_session.commit()

    app.dependency_overrides[get_db] = lambda: db_session
    try:
        client = TestClient(app)
        login_res = client.post("/api/auth/admin/login", json={"email": "admin@printplatform.local", "password": "AdminPass123!"})
        assert login_res.status_code == 200
        token = login_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # 1. Database Status
        db_res = client.get("/api/admin/database/status", headers=headers)
        assert db_res.status_code == 200
        db_data = db_res.json()
        assert db_data["status"] == "CONNECTED"
        assert "latency_ms" in db_data
        assert db_data["total_users"] >= 1
        assert db_data["total_departments"] >= 1

        # 2. Roster Template Download
        tpl_res = client.get("/api/admin/users/import/template", headers=headers)
        assert tpl_res.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in tpl_res.headers["content-type"]
        wb = openpyxl.load_workbook(io.BytesIO(tpl_res.content))
        assert "Student Roster Template" in wb.sheetnames
        ws = wb["Student Roster Template"]
        assert ws.cell(row=1, column=1).value == "Roll Number"

        # 3. Import CSV Roster
        csv_content = (
            "Roll Number,Full Name,Email,Phone,Department Code,Department Name,Role\n"
            "2026-IT-101,Rohan Kumar,rohan.it@campus.edu,9876540001,IT,Information Technology,USER\n"
            "2026-IT-102,Ananya Roy,ananya.it@campus.edu,9876540002,IT,Information Technology,USER\n"
        )
        files = {"file": ("roster.csv", csv_content.encode("utf-8"), "text/csv")}
        import_res = client.post("/api/admin/users/import", headers=headers, files=files)
        assert import_res.status_code == 200
        res_data = import_res.json()
        assert res_data["total_rows"] == 2
        assert res_data["created"] == 2
        assert res_data["failed"] == 0

        # Verify users created in DB
        u1 = db_session.query(User).filter(User.email == "rohan.it@campus.edu").first()
        assert u1 is not None
        assert u1.roll_number == "2026-IT-101"
        assert u1.department == "Information Technology"

        # 4. Import Excel Roster with an update
        wb_import = openpyxl.Workbook()
        ws_import = wb_import.active
        ws_import.append(["Roll Number", "Full Name", "Email", "Phone", "Department Code", "Department Name", "Role"])
        # Update Rohan and add a new user
        ws_import.append(["2026-IT-101", "Rohan Kumar Updated", "rohan.it@campus.edu", "9999999999", "IT", "Information Technology", "USER"])
        ws_import.append(["2026-MECH-201", "Karthik Raja", "karthik.mech@campus.edu", "9876540003", "MECH", "Mechanical Engineering", "USER"])
        excel_buf = io.BytesIO()
        wb_import.save(excel_buf)
        excel_buf.seek(0)

        files_xlsx = {"file": ("roster.xlsx", excel_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        xlsx_res = client.post("/api/admin/users/import", headers=headers, files=files_xlsx)
        assert xlsx_res.status_code == 200
        xlsx_data = xlsx_res.json()
        assert xlsx_data["total_rows"] == 2
        assert xlsx_data["created"] == 1
        assert xlsx_data["updated"] == 1

        u_updated = db_session.query(User).filter(User.email == "rohan.it@campus.edu").first()
        assert u_updated.full_name == "Rohan Kumar Updated"
        assert u_updated.department == "Information Technology"

    finally:
        app.dependency_overrides.clear()
