import io
import pytest
import openpyxl
from fastapi.testclient import TestClient
from app.main import app
from app.db.session import get_db
from app.db.models.user import User, UserRole
from app.db.models.department import Department
from app.db.models.document import Document
from app.db.models.print_server import PrintServer, PrintServerStatus
from app.db.models.printer import Printer
from app.db.models.order import Order, OrderStatus
from app.db.models.print_job import PrintJob, PrintJobStatus
from app.config.security import hash_password

def test_exports_require_admin_auth(db_session):
    app.dependency_overrides[get_db] = lambda: db_session
    try:
        client = TestClient(app)
        assert client.get('/api/reports/export/excel').status_code in (401, 403)
        assert client.get('/api/reports/export/pdf').status_code in (401, 403)
    finally:
        app.dependency_overrides.clear()

def test_generate_excel_and_pdf_exports(db_session):
    admin = User(
        email='admin@printplatform.local',
        password_hash=hash_password('AdminPass123!'),
        role=UserRole.ADMIN,
        is_active=True
    )
    dept = Department(code='MECH_EXP', name='Mechanical Engineering', description='Mech')
    db_session.add(admin)
    db_session.add(dept)
    db_session.flush()

    student = User(
        email='student_export@example.com',
        full_name='Export Test Student',
        password_hash=hash_password('Pass123!'),
        role=UserRole.USER,
        department_id=dept.id,
        department=dept.name,
        roll_number='MECH-2026-01',
        is_active=True
    )
    server = PrintServer(
        id='srv_exp_01',
        name='Station 1',
        device_token_hash='hash123',
        status=PrintServerStatus.ONLINE
    )
    db_session.add(student)
    db_session.add(server)
    db_session.flush()

    printer = Printer(
        id='prt_exp_01',
        server_id=server.id,
        department_id=dept.id,
        cups_printer_name='HP_LaserJet_Mech',
        display_name='HP LaserJet Mech 01',
        supports_color=True,
        supports_duplex=True,
        is_active=True
    )
    doc = Document(
        id='doc_exp_01',
        user_id=student.id,
        original_filename='Lab_Report_Export.pdf',
        stored_filename='doc_exp_01.pdf',
        storage_key='docs/doc_exp_01.pdf',
        sha256='e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855',
        file_size=10240,
        page_count=5,
        mime_type='application/pdf'
    )
    db_session.add(printer)
    db_session.add(doc)
    db_session.flush()

    order = Order(
        id='ord_exp_01',
        user_id=student.id,
        roll_number=student.roll_number,
        department=dept.name,
        document_id=doc.id,
        print_server_id=server.id,
        print_settings={'colour': True, 'sides': 'two-sided', 'copies': 2},
        total_pages=5,
        copies=2,
        amount=50.00,
        status=OrderStatus.COMPLETED
    )
    db_session.add(order)
    db_session.flush()

    job = PrintJob(
        id='job_exp_01',
        order_id=order.id,
        server_id=server.id,
        printer_id=printer.id,
        status=PrintJobStatus.COMPLETED
    )
    db_session.add(job)
    db_session.commit()

    app.dependency_overrides[get_db] = lambda: db_session
    try:
        client = TestClient(app)
        login_res = client.post('/api/auth/admin/login', json={'email': 'admin@printplatform.local', 'password': 'AdminPass123!'})
        assert login_res.status_code == 200
        token = login_res.json()['access_token']
        # A user with order history must remain loadable in the directory.
        users_res = client.get('/api/admin/users', headers={'Authorization': f'Bearer {token}'})
        assert users_res.status_code == 200
        directory_user = next(item for item in users_res.json()['users'] if item['id'] == student.id)
        assert directory_user['total_orders'] == 1
        assert directory_user['total_pages'] == 5
        assert directory_user['total_spent'] == 50.0
        headers = {'Authorization': f'Bearer {token}'}

        # 1. Test Excel Export
        excel_res = client.get('/api/reports/export/excel', headers=headers)
        assert excel_res.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in excel_res.headers['content-type']
        assert len(excel_res.content) > 1000
        # Check ZIP / XLSX magic bytes (PK..)
        assert excel_res.content[:2] == b"PK"

        # Validate that Excel sheet has 4 sheets and contains our seeded job
        wb = openpyxl.load_workbook(io.BytesIO(excel_res.content))
        assert "Executive Summary" in wb.sheetnames
        assert "Departments" in wb.sheetnames
        assert "User Statistics" in wb.sheetnames
        assert "Detailed Print Jobs" in wb.sheetnames

        ws4 = wb["Detailed Print Jobs"]
        assert ws4.max_row >= 2  # Header + at least 1 job row
        row2 = [cell.value for cell in ws4[2]]
        assert row2[0] == "job_exp_01"
        assert row2[1] == "Lab_Report_Export.pdf"
        assert row2[5] == 5  # pages

        # 2. Test PDF Export
        pdf_res = client.get('/api/reports/export/pdf', headers=headers)
        assert pdf_res.status_code == 200
        assert "application/pdf" in pdf_res.headers['content-type']
        assert len(pdf_res.content) > 500
        # Check PDF magic bytes (%PDF)
        assert pdf_res.content[:4] == b"%PDF"

    finally:
        app.dependency_overrides.clear()
