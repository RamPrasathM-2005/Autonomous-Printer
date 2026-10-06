import io
from datetime import datetime, timezone
from typing import Optional
from sqlalchemy.orm import Session

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
)

from app.db.models.department import Department
from app.db.models.user import User
from app.db.models.printer import Printer
from app.services.report_service import report_service

class ExportService:
    @staticmethod
    def generate_excel_report(
        db: Session,
        department_id: Optional[int] = None,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
        user_id: Optional[int] = None,
        printer_id: Optional[str] = None,
        job_status: Optional[str] = None,
        is_color: Optional[bool] = None,
    ) -> bytes:
        report_data = report_service.generate_department_report(
            db=db,
            department_id=department_id,
            start_date=start_date,
            end_date=end_date,
            user_id=user_id,
            printer_id=printer_id,
            job_status=job_status,
            is_color=is_color,
        )

        job_logs = report_service.get_print_job_logs(
            db=db,
            department_id=department_id,
            user_id=user_id,
            printer_id=printer_id,
            job_status=job_status,
            is_color=is_color,
            page=1,
            page_size=2000,
        )

        all_departments = db.query(Department).all()
        s = report_data.summary

        wb = openpyxl.Workbook()

        # Styles
        primary_fill = PatternFill(start_color="1A56DB", end_color="1A56DB", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=16, bold=True, color="1A56DB")
        bold_font = Font(name="Calibri", size=11, bold=True)
        regular_font = Font(name="Calibri", size=11)
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        # ---------------- SHEET 1: Executive Summary ----------------
        ws1 = wb.active
        ws1.title = "Executive Summary"
        ws1.views.sheetView[0].showGridLines = True

        ws1.append(["Achuppori Autonomous Cloud Print - Executive Summary"])
        ws1.cell(row=1, column=1).font = title_font
        ws1.append([f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"])
        ws1.cell(row=2, column=1).font = Font(name="Calibri", size=10, italic=True, color="64748B")
        ws1.append([f"Scope: {s.department_name} ({s.department_code or 'ALL'})"])
        ws1.append([])

        ws1.append(["Key Performance Metric", "Value"])
        ws1.cell(row=5, column=1).fill = primary_fill
        ws1.cell(row=5, column=1).font = header_font
        ws1.cell(row=5, column=2).fill = primary_fill
        ws1.cell(row=5, column=2).font = header_font

        decided = s.successful_jobs + s.failed_jobs
        success_rate = round((s.successful_jobs / decided * 100), 1) if decided > 0 else 100.0

        metrics = [
            ("Total Print Jobs Submitted", s.total_jobs),
            ("Successful Jobs Completed", s.successful_jobs),
            ("Failed / Aborted Jobs", s.failed_jobs),
            ("Platform Success Rate (%)", f"{success_rate}%"),
            ("Total Pages Printed", s.total_pages),
            ("Black & White Pages", s.bw_pages),
            ("Color Pages", s.color_pages),
            ("Average Pages per Job", round(s.avg_pages_per_job, 1)),
            ("Total Amount Billed (₹)", s.total_amount),
        ]

        for idx, (label, val) in enumerate(metrics, start=6):
            ws1.append([label, val])
            cell_label = ws1.cell(row=idx, column=1)
            cell_val = ws1.cell(row=idx, column=2)
            cell_label.font = regular_font
            cell_val.font = bold_font
            cell_label.border = thin_border
            cell_val.border = thin_border
            if idx % 2 == 1:
                cell_label.fill = alt_fill
                cell_val.fill = alt_fill
            if "Amount" in label:
                cell_val.number_format = '"₹"#,##0.00'

        # ---------------- SHEET 2: Departments ----------------
        ws2 = wb.create_sheet(title="Departments")
        ws2.views.sheetView[0].showGridLines = True
        ws2.append(["Department Code", "Department Name", "Users Count", "Printers Assigned", "Description", "Created Date"])
        for col_idx in range(1, 7):
            cell = ws2.cell(row=1, column=col_idx)
            cell.fill = primary_fill
            cell.font = header_font

        for r_idx, d in enumerate(all_departments, start=2):
            users_count = db.query(User).filter((User.department_id == d.id) | (User.department == d.name)).count()
            printers_count = db.query(Printer).filter(Printer.department_id == d.id).count()
            row_data = [
                d.code,
                d.name,
                users_count,
                printers_count,
                d.description or "-",
                d.created_at.strftime("%Y-%m-%d") if d.created_at else "-",
            ]
            ws2.append(row_data)
            for col_idx in range(1, 7):
                c = ws2.cell(row=r_idx, column=col_idx)
                c.font = regular_font
                c.border = thin_border
                if r_idx % 2 == 1:
                    c.fill = alt_fill

        # ---------------- SHEET 3: User Statistics ----------------
        ws3 = wb.create_sheet(title="User Statistics")
        ws3.views.sheetView[0].showGridLines = True
        ws3.append(["Roll Number / ID", "Full Name", "Email", "Jobs Submitted", "Pages Printed", "Color Pages", "B&W Pages", "Success Rate (%)"])
        for col_idx in range(1, 9):
            cell = ws3.cell(row=1, column=col_idx)
            cell.fill = primary_fill
            cell.font = header_font

        for r_idx, u in enumerate(report_data.user_stats, start=2):
            ws3.append([
                u.roll_number or "-",
                u.user_name,
                u.email or "-",
                u.total_jobs,
                u.total_pages,
                u.color_pages,
                u.bw_pages,
                f"{u.success_rate}%",
            ])
            for col_idx in range(1, 9):
                c = ws3.cell(row=r_idx, column=col_idx)
                c.font = regular_font
                c.border = thin_border
                if r_idx % 2 == 1:
                    c.fill = alt_fill

        # ---------------- SHEET 4: Detailed Print Jobs ----------------
        ws4 = wb.create_sheet(title="Detailed Print Jobs")
        ws4.views.sheetView[0].showGridLines = True
        ws4.append(["Job ID", "Document Name", "Requester", "Department", "Printer Queue", "Pages", "Copies", "Color Mode", "Duplex", "Amount (₹)", "Status", "Submitted At", "Completed At"])
        for col_idx in range(1, 14):
            cell = ws4.cell(row=1, column=col_idx)
            cell.fill = primary_fill
            cell.font = header_font

        for r_idx, job in enumerate(job_logs.jobs, start=2):
            ws4.append([
                job.id,
                job.document_name,
                job.user_name or "Guest",
                job.department_name or "General",
                job.printer_name or "-",
                job.pages,
                job.copies,
                "Color" if job.is_color else "B&W",
                "Duplex" if job.is_duplex else "Single",
                job.amount,
                job.status,
                job.created_at,
                job.completed_at or "-",
            ])
            for col_idx in range(1, 14):
                c = ws4.cell(row=r_idx, column=col_idx)
                c.font = regular_font
                c.border = thin_border
                if r_idx % 2 == 1:
                    c.fill = alt_fill
                if col_idx == 10:
                    c.number_format = '"₹"#,##0.00'

        # Auto-fit columns across all sheets
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    if "\n" in val_str:
                        val_str = val_str.split("\n")[0]
                    max_len = max(max_len, len(val_str))
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def generate_pdf_report(
        db: Session,
        department_id: Optional[int] = None,
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
        user_id: Optional[int] = None,
        printer_id: Optional[str] = None,
        job_status: Optional[str] = None,
        is_color: Optional[bool] = None,
    ) -> bytes:
        report_data = report_service.generate_department_report(
            db=db,
            department_id=department_id,
            start_date=start_date,
            end_date=end_date,
            user_id=user_id,
            printer_id=printer_id,
            job_status=job_status,
            is_color=is_color,
        )

        all_departments = db.query(Department).all()
        s = report_data.summary

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1A56DB"),
            fontName="Helvetica-Bold",
        )
        subtitle_style = ParagraphStyle(
            "DocSubTitle",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#64748B"),
            fontName="Helvetica",
        )
        h2_style = ParagraphStyle(
            "SectionH2",
            parent=styles["Heading2"],
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#0D1117"),
            fontName="Helvetica-Bold",
            spaceAfter=6,
        )
        cell_style = ParagraphStyle(
            "CellNormal",
            parent=styles["Normal"],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#1E293B"),
            fontName="Helvetica",
        )
        cell_bold = ParagraphStyle(
            "CellBold",
            parent=styles["Normal"],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#0D1117"),
            fontName="Helvetica-Bold",
        )
        header_cell_style = ParagraphStyle(
            "HeaderCell",
            parent=styles["Normal"],
            fontSize=8.5,
            leading=11,
            textColor=colors.white,
            fontName="Helvetica-Bold",
        )

        elements = []

        # Header Banner
        elements.append(Paragraph("ACHUPPORI AUTONOMOUS CLOUD PRINT", title_style))
        elements.append(Paragraph("Official Administrative Printing & Utilization Report", subtitle_style))
        gen_time = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        elements.append(Paragraph(f"Generated: {gen_time} &bull; Scope: {s.department_name} ({s.department_code or 'ALL'})", subtitle_style))
        elements.append(Spacer(1, 10))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1A56DB"), spaceBefore=0, spaceAfter=14))

        # Executive Metrics Grid
        elements.append(Paragraph("Executive Performance Metrics", h2_style))
        decided = s.successful_jobs + s.failed_jobs
        success_rate = round((s.successful_jobs / decided * 100), 1) if decided > 0 else 100.0

        kpi_data = [
            [
                Paragraph(f"<b>Total Jobs:</b> {s.total_jobs}", cell_style),
                Paragraph(f"<b>Successful:</b> {s.successful_jobs}", cell_style),
                Paragraph(f"<b>Failed:</b> {s.failed_jobs}", cell_style),
                Paragraph(f"<b>Success Rate:</b> {success_rate}%", cell_style),
            ],
            [
                Paragraph(f"<b>Pages Printed:</b> {s.total_pages}", cell_style),
                Paragraph(f"<b>B&W Pages:</b> {s.bw_pages}", cell_style),
                Paragraph(f"<b>Color Pages:</b> {s.color_pages}", cell_style),
                Paragraph(f"<b>Total Revenue:</b> ₹{s.total_amount:.2f}", cell_bold),
            ],
        ]
        kpi_table = Table(kpi_data, colWidths=[130, 130, 130, 130])
        kpi_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F0F5FF")),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#BFCFEF")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDE1E7")),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 16))

        # Department Allocation Table
        elements.append(Paragraph("Department Directory & Allocations", h2_style))
        dept_rows = [
            [
                Paragraph("Code", header_cell_style),
                Paragraph("Department Name", header_cell_style),
                Paragraph("Users", header_cell_style),
                Paragraph("Printers", header_cell_style),
                Paragraph("Description", header_cell_style),
            ]
        ]
        for d in all_departments:
            users_count = db.query(User).filter((User.department_id == d.id) | (User.department == d.name)).count()
            printers_count = db.query(Printer).filter(Printer.department_id == d.id).count()
            dept_rows.append([
                Paragraph(d.code, cell_bold),
                Paragraph(d.name, cell_style),
                Paragraph(str(users_count), cell_style),
                Paragraph(str(printers_count), cell_style),
                Paragraph(d.description or "-", cell_style),
            ])

        dept_table = Table(dept_rows, colWidths=[80, 200, 70, 70, 100])
        dept_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A56DB")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(dept_table)
        elements.append(Spacer(1, 16))

        # Top Users Table
        if report_data.user_stats:
            elements.append(Paragraph("Active Print Users", h2_style))
            user_rows = [
                [
                    Paragraph("Roll / ID", header_cell_style),
                    Paragraph("Name", header_cell_style),
                    Paragraph("Email", header_cell_style),
                    Paragraph("Jobs", header_cell_style),
                    Paragraph("Pages", header_cell_style),
                    Paragraph("Success Rate", header_cell_style),
                ]
            ]
            for u in report_data.user_stats[:12]:
                user_rows.append([
                    Paragraph(u.roll_number or "-", cell_bold),
                    Paragraph(u.user_name, cell_style),
                    Paragraph(u.email or "-", cell_style),
                    Paragraph(str(u.total_jobs), cell_style),
                    Paragraph(str(u.total_pages), cell_style),
                    Paragraph(f"{u.success_rate}%", cell_style),
                ])

            user_table = Table(user_rows, colWidths=[80, 130, 150, 50, 50, 60])
            user_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ]))
            elements.append(user_table)

        doc.build(elements)
        return buffer.getvalue()

export_service = ExportService()
