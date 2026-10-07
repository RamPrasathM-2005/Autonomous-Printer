import io
import csv
import time
import secrets
from typing import List, Dict, Any, Optional
from sqlalchemy.orm import Session
from sqlalchemy import func, text
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.db.models.department import Department
from app.db.models.user import User, UserRole
from app.db.models.printer import Printer
from app.db.models.print_server import PrintServer, PrintServerStatus
from app.db.models.discovered_printer import DiscoveredPrinter
from app.db.models.print_job import PrintJob, PrintJobStatus
from app.db.models.order import Order
from app.db.models.document import Document
from app.config.security import hash_password
from app.schemas.admin import (
    AdminDashboardResponse,
    RecentPrintJob,
    DepartmentSummary,
    AdminPrinterItem,
    AdminPrinterCreate,
    AdminPrinterUpdate,
    PrinterTestConnectionResponse,
    AdminDiscoveredPrinterItem,
    AdminPrintAgentItem,
    AdminPrintAgentCreate,
    AdminPrintServerUpdate,
    AdminAgentCupsPrinterItem,
    AdminUserItem,
    AdminUserCreate,
    AdminUserUpdate,
    PaginatedAdminUsersResponse,
    AdminSystemSettings,
    AdminSystemSettingsUpdate,
    AdminDatabaseStatusResponse,
    UserImportSummary,
)
from app.utils.errors import AppException


class AdminService:
    @staticmethod
    def get_dashboard_metrics(db: Session) -> AdminDashboardResponse:
        total_departments = db.query(Department).count()
        total_users = db.query(User).count()
        total_printers = db.query(Printer).count()
        
        # Load all print jobs with linked order and document
        query_results = (
            db.query(PrintJob, Order, Document, User, Printer, Department)
            .outerjoin(Order, Order.id == PrintJob.order_id)
            .outerjoin(Document, Document.id == Order.document_id)
            .outerjoin(User, User.id == Order.user_id)
            .outerjoin(Printer, Printer.id == PrintJob.printer_id)
            .outerjoin(Department, (Department.id == User.department_id) | (Department.name == Order.department) | (Department.code == Order.department))
            .order_by(PrintJob.created_at.desc())
            .all()
        )
        
        total_print_jobs = len(query_results)
        successful_print_jobs = 0
        failed_print_jobs = 0
        total_pages_printed = 0
        color_pages = 0
        bw_pages = 0
        color_jobs = 0
        bw_jobs = 0
        
        recent_activity: List[RecentPrintJob] = []
        
        # Maps for department aggregation
        dept_jobs_map: Dict[int, int] = {}
        dept_pages_map: Dict[int, int] = {}
        
        for idx, (job, order, doc, user, printer, dept) in enumerate(query_results):
            # Status check
            status_val = job.status.value if hasattr(job.status, "value") else str(job.status)
            is_success = status_val == PrintJobStatus.COMPLETED.value
            is_failed = status_val in [
                PrintJobStatus.FAILED.value,
                PrintJobStatus.FINAL_FAILED.value,
                PrintJobStatus.CANCELLED.value
            ]
            
            if is_success:
                successful_print_jobs += 1
            elif is_failed:
                failed_print_jobs += 1
                
            # Pages & Color check
            copies = order.copies if (order and order.copies) else 1
            pages_per_copy = (order.total_pages if order and order.total_pages else (doc.page_count if doc and doc.page_count else 1))
            job_pages = pages_per_copy * copies
            
            is_color = False
            if order and order.print_settings:
                c_val = order.print_settings.get("colour", order.print_settings.get("color", False))
                if isinstance(c_val, str):
                    is_color = c_val.lower() in ["true", "color", "colour", "1"]
                else:
                    is_color = bool(c_val)
                    
            if is_color:
                color_jobs += 1
            else:
                bw_jobs += 1
                
            if is_success:
                total_pages_printed += job_pages
                if is_color:
                    color_pages += job_pages
                else:
                    bw_pages += job_pages
                    
            # Department tracking
            dept_id = dept.id if dept else (user.department_id if user and user.department_id else None)
            if dept_id:
                dept_jobs_map[dept_id] = dept_jobs_map.get(dept_id, 0) + 1
                if is_success:
                    dept_pages_map[dept_id] = dept_pages_map.get(dept_id, 0) + job_pages
                    
            # Recent items (limit to 15)
            if idx < 15:
                # Find printer display name
                printer_name = "Default Printer"
                if printer and printer.display_name:
                    printer_name = printer.display_name
                elif order and order.print_settings:
                    printer_name = order.print_settings.get("printer_name") or order.print_settings.get("cups_printer_name") or "Local Station Printer"
                    
                recent_activity.append(RecentPrintJob(
                    id=job.id,
                    order_id=order.id if order else "",
                    document_name=doc.original_filename if doc else (order.id if order else "Document"),
                    user_name=(user.full_name or user.email) if user else "Walk-in Guest",
                    user_email=user.email if user else None,
                    department_name=dept.name if dept else (order.department if order and order.department else (user.department if user and user.department else "General")),
                    printer_name=printer_name,
                    pages=job_pages,
                    copies=copies,
                    is_color=is_color,
                    status=status_val,
                    created_at=job.created_at.strftime("%Y-%m-%d %H:%M:%S") if job.created_at else "",
                    completed_at=job.completed_at.strftime("%Y-%m-%d %H:%M:%S") if job.completed_at else None,
                    error_message=job.error_message
                ))
                
        # Calculate success rate
        decided_jobs = successful_print_jobs + failed_print_jobs
        success_rate = round((successful_print_jobs / decided_jobs * 100), 1) if decided_jobs > 0 else 100.0
        
        # Build department summaries
        all_departments = db.query(Department).all()
        department_summaries: List[DepartmentSummary] = []
        for d in all_departments:
            users_count = db.query(User).filter(
                (User.department_id == d.id) | (User.department == d.name) | (User.department == d.code)
            ).count()
            printers_count = db.query(Printer).filter(Printer.department_id == d.id).count()
            department_summaries.append(DepartmentSummary(
                id=d.id,
                code=d.code,
                name=d.name,
                users_count=users_count,
                printers_count=printers_count,
                jobs_count=dept_jobs_map.get(d.id, 0),
                pages_count=dept_pages_map.get(d.id, 0),
            ))
            
        return AdminDashboardResponse(
            total_departments=total_departments,
            total_users=total_users,
            total_printers=total_printers,
            total_print_jobs=total_print_jobs,
            total_pages_printed=total_pages_printed,
            successful_print_jobs=successful_print_jobs,
            failed_print_jobs=failed_print_jobs,
            color_pages=color_pages,
            bw_pages=bw_pages,
            color_jobs=color_jobs,
            bw_jobs=bw_jobs,
            success_rate=success_rate,
            recent_activity=recent_activity,
            department_summary=department_summaries
        )

    @staticmethod
    def get_printers(db: Session) -> List[AdminPrinterItem]:
        from datetime import datetime, timezone, timedelta
        from app.db.models.print_server import PrintServer, PrintServerStatus
        from app.config.settings import settings

        now = datetime.now(timezone.utc)
        timeout = timedelta(seconds=settings.HEARTBEAT_TIMEOUT_SECONDS)

        printers = db.query(Printer).all()
        results: List[AdminPrinterItem] = []

        for p in printers:
            server = db.query(PrintServer).filter(PrintServer.id == p.server_id).first()
            dept = db.query(Department).filter(Department.id == p.department_id).first() if p.department_id else None
            job_count = db.query(PrintJob).filter(PrintJob.printer_id == p.id).count()

            # Determine dynamic server status
            server_status = server.status.value if (server and hasattr(server.status, "value")) else ("OFFLINE" if not server else str(server.status))
            if server and server.status not in [PrintServerStatus.DISABLED, PrintServerStatus.MAINTENANCE]:
                hb = server.last_heartbeat
                if hb:
                    if hb.tzinfo is None:
                        hb = hb.replace(tzinfo=timezone.utc)
                    if hb < now - timeout:
                        server_status = "OFFLINE"
                else:
                    server_status = "OFFLINE"

            results.append(AdminPrinterItem(
                id=p.id,
                server_id=p.server_id,
                server_name=server.name if server else None,
                server_status=server_status,
                server_location=server.location if server else None,
                department_id=p.department_id,
                department_name=dept.name if dept else None,
                cups_printer_name=p.cups_printer_name,
                display_name=p.display_name,
                ip_address=p.ip_address,
                protocol=p.protocol or "Socket",
                device_uri=p.device_uri,
                model=p.model,
                location=p.location or (server.location if server else None),
                description=p.description,
                supports_color=p.supports_color,
                supports_duplex=p.supports_duplex,
                is_enabled=p.is_enabled if p.is_enabled is not None else True,
                is_active=p.is_active,
                printer_state=p.printer_status or (server.printer_state if server else "UNKNOWN"),
                paper_state=server.paper_state if server else "UNKNOWN",
                last_heartbeat=server.last_heartbeat.strftime("%Y-%m-%d %H:%M:%S") if (server and server.last_heartbeat) else None,
                total_jobs_count=job_count,
                created_at=p.created_at.strftime("%Y-%m-%d %H:%M:%S") if p.created_at else "",
                last_tested_at=p.last_tested_at.strftime("%Y-%m-%d %H:%M:%S") if p.last_tested_at else None,
                test_status=p.test_status or "PENDING",
                test_message=p.test_message,
                has_mismatch=p.has_mismatch,
                mismatch_details=p.mismatch_details,
            ))

        return results

    @staticmethod
    def create_printer(data: AdminPrinterCreate, db: Session) -> AdminPrinterItem:
        import uuid
        from datetime import datetime, timezone

        server = db.query(PrintServer).filter(PrintServer.id == data.server_id).first()
        if not server:
            raise AppException(f"Selected Print Agent '{data.server_id}' does not exist.", status_code=400)
        if server.status == PrintServerStatus.DISABLED:
            raise AppException("Selected Print Agent is currently disabled.", status_code=400)

        if data.department_id:
            dept = db.query(Department).filter(Department.id == data.department_id).first()
            if not dept:
                raise AppException("Selected department does not exist.", status_code=400)

        # Check for duplicate CUPS name on same agent
        existing = db.query(Printer).filter(
            Printer.server_id == data.server_id,
            Printer.cups_printer_name == data.cups_printer_name.strip()
        ).first()
        if existing:
            raise AppException(
                f"Printer '{data.cups_printer_name}' is already registered on Print Agent '{data.server_id}'.",
                status_code=400
            )

        # Compute device_uri if not explicitly provided
        device_uri = (data.device_uri or "").strip()
        proto = (data.protocol or "Socket").strip()
        ip = (data.ip_address or "").strip()
        if not device_uri and ip:
            if proto.lower() == "ipp":
                device_uri = f"ipp://{ip}:631/ipp/print"
            elif proto.lower() == "lpd":
                device_uri = f"lpd://{ip}/queue"
            else:
                device_uri = f"socket://{ip}:9100"

        # Unique printer ID
        printer_id = f"PRN-{uuid.uuid4().hex[:8].upper()}"

        dept_id = data.department_id if data.department_id is not None else server.department_id

        new_printer = Printer(
            id=printer_id,
            server_id=data.server_id,
            department_id=dept_id,
            cups_printer_name=data.cups_printer_name.strip(),
            display_name=data.display_name.strip(),
            ip_address=ip or None,
            protocol=proto,
            device_uri=device_uri or None,
            model=data.model.strip() if data.model else None,
            location=data.location.strip() if data.location else None,
            description=data.description.strip() if data.description else None,
            supports_color=data.supports_color,
            supports_duplex=data.supports_duplex,
            is_enabled=data.is_enabled,
            is_active=data.is_active if data.is_active is not None else False,
            printer_status="READY",
            test_status="VERIFIED" if (data.is_active is True) else "PENDING",
            test_message="Registered and activated by administrator." if (data.is_active is True) else "Awaiting test connection verification.",
            has_mismatch=False,
            mismatch_details=None,
        )

        db.add(new_printer)

        # If it was in DiscoveredPrinter, remove from discovered queue
        disc_id = f"{data.server_id}::{data.cups_printer_name.strip()}"
        disc = db.query(DiscoveredPrinter).filter(DiscoveredPrinter.id == disc_id).first()
        if disc:
            db.delete(disc)

        db.commit()
        db.refresh(new_printer)

        printers = AdminService.get_printers(db)
        for p in printers:
            if p.id == new_printer.id:
                return p
        raise AppException("Failed to load newly created printer", status_code=500)

    @staticmethod
    def test_printer_connection(printer_id: str, db: Session) -> PrinterTestConnectionResponse:
        import socket
        import requests
        from datetime import datetime, timezone
        from app.config.settings import settings

        printer = db.query(Printer).filter(Printer.id == printer_id).first()
        if not printer:
            raise AppException(f"Printer '{printer_id}' not found.", status_code=404)

        server = db.query(PrintServer).filter(PrintServer.id == printer.server_id).first()
        if not server:
            printer.is_active = False
            printer.test_status = "FAILED"
            printer.test_message = f"Assigned Print Agent '{printer.server_id}' does not exist."
            printer.last_tested_at = datetime.now(timezone.utc)
            db.commit()
            return PrinterTestConnectionResponse(
                success=False,
                printer_id=printer.id,
                is_active=False,
                cups_exists=False,
                reachable=False,
                available=False,
                status="FAILED",
                message=printer.test_message,
                tested_at=printer.last_tested_at.strftime("%Y-%m-%d %H:%M:%S")
            )

        if server.status == PrintServerStatus.DISABLED:
            printer.is_active = False
            printer.test_status = "FAILED"
            printer.test_message = f"Assigned Print Agent '{server.name}' is currently disabled."
            printer.last_tested_at = datetime.now(timezone.utc)
            db.commit()
            return PrinterTestConnectionResponse(
                success=False,
                printer_id=printer.id,
                is_active=False,
                cups_exists=False,
                reachable=False,
                available=False,
                status="FAILED",
                message=printer.test_message,
                tested_at=printer.last_tested_at.strftime("%Y-%m-%d %H:%M:%S")
            )

        # 1. Attempt to instruct assigned Print Agent via HTTP endpoint
        agent_responded = False
        agent_result = {}
        dispatch_headers = {
            "Authorization": f"Bearer {settings.INTERNAL_AGENT_TOKEN}",
            "X-Internal-Token": settings.INTERNAL_AGENT_TOKEN
        }
        test_payload = {
            "cups_printer_name": printer.cups_printer_name,
            "ip_address": printer.ip_address,
            "device_uri": printer.device_uri,
            "protocol": printer.protocol or "Socket"
        }

        candidate_urls = []
        if server.ip_address:
            candidate_urls.extend([
                f"http://{server.ip_address}:5001/local/test-printer",
                f"http://{server.ip_address}:5000/local/test-printer"
            ])
        candidate_urls.extend([
            "http://127.0.0.1:5001/local/test-printer",
            "http://127.0.0.1:5000/local/test-printer"
        ])

        for url in candidate_urls:
            try:
                resp = requests.post(url, json=test_payload, headers=dispatch_headers, timeout=3)
                if resp.status_code == 200:
                    agent_result = resp.json()
                    agent_responded = True
                    break
            except Exception:
                continue

        # 2. Evaluate connection test
        if agent_responded:
            success = bool(agent_result.get("success", False))
            cups_exists = bool(agent_result.get("cups_exists", False))
            reachable = bool(agent_result.get("reachable", False))
            available = bool(agent_result.get("available", False))
            message = str(agent_result.get("message") or "")
        else:
            # Fallback probe if print agent HTTP endpoint was unreachable:
            port = 9100
            proto_lower = (printer.protocol or "").lower()
            if "ipp" in proto_lower or (printer.device_uri and "631" in printer.device_uri):
                port = 631
            elif "lpd" in proto_lower or (printer.device_uri and "515" in printer.device_uri):
                port = 515

            target_ip = (printer.ip_address or "").strip()
            if target_ip:
                try:
                    s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
                    s.settimeout(2.0)
                    s.connect((target_ip, port))
                    s.close()
                    reachable = True
                    cups_exists = True
                    available = True
                    success = True
                    message = f"Network probe verified connection to {target_ip}:{port} ({printer.protocol})."
                except Exception as ex:
                    reachable = False
                    cups_exists = False
                    available = False
                    success = False
                    message = f"Connection failed to {target_ip}:{port} ({ex}). Printer is unreachable or offline."
            else:
                reachable = False
                cups_exists = False
                available = False
                success = False
                message = "Assigned Print Agent is offline and no static IP address is configured for direct probe."

        # 3. Only after a successful validation should the printer be marked as Active
        now = datetime.now(timezone.utc)
        printer.last_tested_at = now
        if success:
            printer.is_active = True
            printer.test_status = "SUCCESS"
            printer.printer_status = "READY"
            printer.test_message = message
            printer.has_mismatch = False
            printer.mismatch_details = None
        else:
            printer.is_active = False
            printer.test_status = "FAILED"
            printer.printer_status = "OFFLINE"
            printer.test_message = message

        db.commit()
        db.refresh(printer)

        return PrinterTestConnectionResponse(
            success=success,
            printer_id=printer.id,
            is_active=printer.is_active,
            cups_exists=cups_exists,
            reachable=reachable,
            available=available,
            status=printer.test_status,
            message=printer.test_message or "",
            tested_at=printer.last_tested_at.strftime("%Y-%m-%d %H:%M:%S")
        )

    @staticmethod
    def update_printer(printer_id: str, data: AdminPrinterUpdate, db: Session) -> AdminPrinterItem:
        printer = db.query(Printer).filter(Printer.id == printer_id).first()
        if not printer:
            raise AppException(f"Printer '{printer_id}' not found", status_code=404)

        if data.display_name is not None:
            printer.display_name = data.display_name.strip()
        if data.cups_printer_name is not None:
            printer.cups_printer_name = data.cups_printer_name.strip()
        if data.ip_address is not None:
            printer.ip_address = data.ip_address.strip() or None
        if data.protocol is not None:
            printer.protocol = data.protocol.strip()
        if data.device_uri is not None:
            printer.device_uri = data.device_uri.strip() or None
        if data.model is not None:
            printer.model = data.model.strip() or None
        if data.location is not None:
            printer.location = data.location.strip() or None
        if data.description is not None:
            printer.description = data.description.strip() or None
        if data.supports_color is not None:
            printer.supports_color = data.supports_color
        if data.supports_duplex is not None:
            printer.supports_duplex = data.supports_duplex
        if data.is_enabled is not None:
            printer.is_enabled = data.is_enabled
        if data.server_id is not None:
            server = db.query(PrintServer).filter(PrintServer.id == data.server_id).first()
            if not server:
                raise AppException(f"Selected Print Agent '{data.server_id}' not found", status_code=400)
            printer.server_id = data.server_id
        if data.department_id is not None:
            if data.department_id > 0:
                dept = db.query(Department).filter(Department.id == data.department_id).first()
                if not dept:
                    raise AppException("Selected department does not exist", status_code=400)
                printer.department_id = data.department_id
            else:
                printer.department_id = None
        if data.is_active is not None:
            printer.is_active = data.is_active

        db.commit()
        db.refresh(printer)

        printers = AdminService.get_printers(db)
        for p in printers:
            if p.id == printer.id:
                return p
        raise AppException("Failed to load updated printer", status_code=500)

    @staticmethod
    def delete_printer(printer_id: str, db: Session) -> bool:
        printer = db.query(Printer).filter(Printer.id == printer_id).first()
        if not printer:
            raise AppException(f"Printer '{printer_id}' not found", status_code=404)

        # Prevent deletion if active jobs are printing
        active_jobs = db.query(PrintJob).filter(
            PrintJob.printer_id == printer_id,
            PrintJob.status.in_([PrintJobStatus.PRINTING, PrintJobStatus.PROCESSING])
        ).count()
        if active_jobs > 0:
            raise AppException("Cannot delete printer while active print jobs are in progress.", status_code=400)

        db.delete(printer)
        db.commit()
        return True

    @staticmethod
    def get_print_servers(db: Session) -> List[AdminPrintAgentItem]:
        from datetime import datetime, timezone, timedelta
        from app.config.settings import settings

        now = datetime.now(timezone.utc)
        timeout = timedelta(seconds=settings.HEARTBEAT_TIMEOUT_SECONDS)

        servers = db.query(PrintServer).all()
        results: List[AdminPrintAgentItem] = []

        for s in servers:
            dept = db.query(Department).filter(Department.id == s.department_id).first() if s.department_id else None
            p_count = db.query(Printer).filter(Printer.server_id == s.id).count()

            server_status = s.status.value if (s and hasattr(s.status, "value")) else ("OFFLINE" if not s else str(s.status))
            if s.status not in [PrintServerStatus.DISABLED, PrintServerStatus.MAINTENANCE]:
                hb = s.last_heartbeat
                if hb:
                    if hb.tzinfo is None:
                        hb = hb.replace(tzinfo=timezone.utc)
                    if hb < now - timeout:
                        server_status = "OFFLINE"
                else:
                    server_status = "OFFLINE"

            results.append(AdminPrintAgentItem(
                id=s.id,
                name=s.name,
                location=s.location,
                department_id=s.department_id,
                department_name=dept.name if dept else None,
                hostname=s.hostname,
                ip_address=s.ip_address,
                software_version=s.software_version,
                status=server_status,
                is_enabled=s.is_enabled if s.is_enabled is not None else True,
                last_heartbeat=s.last_heartbeat.strftime("%Y-%m-%d %H:%M:%S") if s.last_heartbeat else None,
                printer_count=p_count,
                created_at=s.created_at.strftime("%Y-%m-%d %H:%M:%S") if s.created_at else "",
            ))

        return results

    @staticmethod
    def create_print_server(data: AdminPrintAgentCreate, db: Session) -> AdminPrintAgentItem:
        import hashlib
        import uuid

        dept = None
        if data.department_id:
            dept = db.query(Department).filter(Department.id == data.department_id).first()

        agent_id = (data.id or "").strip()
        if not agent_id:
            dept_code = dept.code if dept else "STATION"
            agent_id = f"AGENT-{dept_code}-{uuid.uuid4().hex[:6].upper()}"

        existing = db.query(PrintServer).filter(PrintServer.id == agent_id).first()
        if existing:
            raise AppException(f"Print Agent with ID '{agent_id}' already exists", status_code=400)

        token = data.token or f"agent-sec-{secrets.token_hex(16)}"
        token_hash = hashlib.sha256(token.encode()).hexdigest()

        location = data.location.strip() if data.location else (f"{dept.name} Station" if dept else data.name.strip())

        server = PrintServer(
            id=agent_id,
            name=data.name.strip(),
            location=location,
            device_token_hash=token_hash,
            department_id=data.department_id,
            hostname=data.hostname.strip() if data.hostname else None,
            ip_address=data.ip_address.strip() if data.ip_address else None,
            software_version=data.software_version.strip() if data.software_version else "1.0.0",
            is_enabled=True,
            status=PrintServerStatus.OFFLINE,
        )
        db.add(server)
        db.commit()
        db.refresh(server)

        agents = AdminService.get_print_servers(db)
        for a in agents:
            if a.id == server.id:
                return a
        raise AppException("Failed to load created print agent", status_code=500)

    @staticmethod
    def get_agent_cups_printers(server_id: str, db: Session) -> List[AdminAgentCupsPrinterItem]:
        import requests
        from app.config.settings import settings

        server = db.query(PrintServer).filter(PrintServer.id == server_id).first()
        if not server:
            raise AppException(f"Print Agent '{server_id}' not found", status_code=404)

        registered_cups = {
            p.cups_printer_name.strip().lower()
            for p in db.query(Printer).filter(Printer.server_id == server_id).all()
        }

        agent_cups: List[dict] = []
        dispatch_headers = {
            "Authorization": f"Bearer {settings.INTERNAL_AGENT_TOKEN}",
            "X-Internal-Token": settings.INTERNAL_AGENT_TOKEN
        }

        urls = []
        if server.ip_address:
            clean_ip = server.ip_address.strip().replace("http://", "").replace("https://", "").split("/")[0].split(":")[0]
            urls.extend([
                f"http://{clean_ip}:5001/local/cups-printers",
                f"http://{clean_ip}:5000/local/cups-printers",
                f"http://{clean_ip}:5001/local/detected-printers",
            ])
        urls.extend([
            "http://127.0.0.1:5001/local/cups-printers",
            "http://127.0.0.1:5000/local/cups-printers",
        ])

        for url in urls:
            try:
                resp = requests.get(url, headers=dispatch_headers, timeout=3)
                if resp.status_code == 200:
                    data = resp.json()
                    if isinstance(data, list):
                        agent_cups = data
                        break
                    elif isinstance(data, dict) and "printers" in data:
                        agent_cups = data["printers"]
                        break
            except Exception:
                continue

        # If direct HTTP request couldn't connect, fall back to DiscoveredPrinter records from heartbeats
        if not agent_cups:
            disc_printers = db.query(DiscoveredPrinter).filter(DiscoveredPrinter.server_id == server_id).all()
            for d in disc_printers:
                agent_cups.append({
                    "cups_printer_name": d.cups_printer_name,
                    "device_uri": d.device_uri,
                    "ip_address": d.ip_address,
                    "status": d.reported_status or "READY"
                })

        results: List[AdminAgentCupsPrinterItem] = []
        seen = set()
        for item in agent_cups:
            c_name = item.get("cups_printer_name") or item.get("name") or ""
            if not c_name or c_name.lower() in seen:
                continue
            seen.add(c_name.lower())
            is_mapped = c_name.lower() in registered_cups
            disp_name = c_name.replace("_", " ")
            results.append(AdminAgentCupsPrinterItem(
                cups_printer_name=c_name,
                display_name=disp_name,
                device_uri=item.get("device_uri"),
                ip_address=item.get("ip_address") or server.ip_address,
                status=item.get("status") or "READY",
                is_mapped=is_mapped
            ))

        return results

    @staticmethod
    def delete_print_server(server_id: str, db: Session) -> bool:
        server = db.query(PrintServer).filter(PrintServer.id == server_id).first()
        if not server:
            raise AppException(f"Print Agent '{server_id}' not found", status_code=404)

        # Automatically cascade delete or unmap assigned printers
        printers = db.query(Printer).filter(Printer.server_id == server_id).all()
        for p in printers:
            db.delete(p)

        db.delete(server)
        db.commit()
        return True

    @staticmethod
    def get_discovered_printers(db: Session) -> List[AdminDiscoveredPrinterItem]:
        discovered = db.query(DiscoveredPrinter).order_by(DiscoveredPrinter.last_seen.desc()).all()
        results: List[AdminDiscoveredPrinterItem] = []
        for d in discovered:
            srv = db.query(PrintServer).filter(PrintServer.id == d.server_id).first()
            results.append(AdminDiscoveredPrinterItem(
                id=d.id,
                server_id=d.server_id,
                server_name=srv.name if srv else None,
                cups_printer_name=d.cups_printer_name,
                ip_address=d.ip_address,
                device_uri=d.device_uri,
                reported_status=d.reported_status or "READY",
                reported_jobs=d.reported_jobs or "0",
                first_seen=d.first_seen.strftime("%Y-%m-%d %H:%M:%S") if d.first_seen else "",
                last_seen=d.last_seen.strftime("%Y-%m-%d %H:%M:%S") if d.last_seen else "",
            ))
        return results

    @staticmethod
    def dismiss_discovered_printer(discovered_id: str, db: Session) -> bool:
        disc = db.query(DiscoveredPrinter).filter(DiscoveredPrinter.id == discovered_id).first()
        if disc:
            db.delete(disc)
            db.commit()
        return True

    @staticmethod
    def update_print_server(server_id: str, data: AdminPrintServerUpdate, db: Session) -> bool:
        from app.db.models.print_server import PrintServer, PrintServerStatus
        server = db.query(PrintServer).filter(PrintServer.id == server_id).first()
        if not server:
            raise AppException(f"Print Server '{server_id}' not found", status_code=404)

        if data.name is not None:
            server.name = data.name.strip()
        if data.status is not None:
            try:
                server.status = PrintServerStatus(data.status.upper())
            except ValueError:
                raise AppException(f"Invalid server status '{data.status}'", status_code=400)
        if data.location is not None:
            server.location = data.location.strip()
        if data.department_id is not None:
            server.department_id = data.department_id if data.department_id > 0 else None
        if data.hostname is not None:
            server.hostname = data.hostname.strip() or None
        if data.ip_address is not None:
            server.ip_address = data.ip_address.strip() or None
        if data.software_version is not None:
            server.software_version = data.software_version.strip() or None
        if data.is_enabled is not None:
            server.is_enabled = data.is_enabled

        db.commit()
        return True

        return True

    @staticmethod
    def get_users(
        db: Session,
        page: int = 1,
        limit: int = 15,
        search: Optional[str] = None,
        department_id: Optional[int] = None,
        role: Optional[str] = None,
        is_active: Optional[bool] = None,
    ) -> PaginatedAdminUsersResponse:
        import math
        query = db.query(User)

        if search:
            term = f"%{search.strip()}%"
            query = query.filter(
                (User.full_name.ilike(term)) |
                (User.email.ilike(term)) |
                (User.roll_number.ilike(term))
            )

        if department_id:
            query = query.filter(User.department_id == department_id)

        if role:
            query = query.filter(User.role == role)

        if is_active is not None:
            query = query.filter(User.is_active == is_active)

        total = query.count()
        total_pages = max(1, math.ceil(total / limit)) if limit > 0 else 1

        users = query.order_by(User.id.desc()).offset((page - 1) * limit).limit(limit).all()

        user_items: List[AdminUserItem] = []
        for u in users:
            dept = db.query(Department).filter(Department.id == u.department_id).first() if u.department_id else None
            orders = db.query(Order).filter(Order.user_id == u.id).all()
            total_orders = len(orders)
            total_pages_sum = sum(o.total_pages or 0 for o in orders)
            total_spent = sum(float(o.total_amount or 0.0) for o in orders)

            user_items.append(AdminUserItem(
                id=u.id,
                email=u.email,
                phone=getattr(u, "phone", None),
                full_name=u.full_name,
                roll_number=u.roll_number,
                department_id=u.department_id,
                department_name=dept.name if dept else (u.department or None),
                role=u.role.value if hasattr(u.role, "value") else str(u.role),
                is_active=u.is_active,
                total_orders=total_orders,
                total_pages=total_pages_sum,
                total_spent=round(total_spent, 2),
                created_at=u.created_at.strftime("%Y-%m-%d %H:%M:%S") if u.created_at else "",
            ))

        return PaginatedAdminUsersResponse(
            users=user_items,
            total=total,
            page=page,
            limit=limit,
            total_pages=total_pages,
        )

    @staticmethod
    def update_user(user_id: int, data: AdminUserUpdate, db: Session) -> AdminUserItem:
        from app.db.models.user import UserRole
        user = db.query(User).filter(User.id == user_id).first()
        if not user:
            raise AppException(f"User with ID {user_id} not found", status_code=404)

        if data.full_name is not None:
            user.full_name = data.full_name.strip()

        if data.email is not None:
            clean_email = data.email.strip().lower()
            if clean_email and clean_email != (user.email or "").lower():
                existing_email = db.query(User).filter(User.email == clean_email, User.id != user.id).first()
                if existing_email:
                    raise AppException("A user with this email address already exists.", status_code=400)
                user.email = clean_email

        if data.roll_number is not None:
            clean_roll = data.roll_number.strip() if data.roll_number.strip() else None
            if clean_roll and clean_roll != user.roll_number:
                existing_roll = db.query(User).filter(User.roll_number == clean_roll, User.id != user.id).first()
                if existing_roll:
                    raise AppException(f"A user with roll number '{clean_roll}' already exists.", status_code=400)
            user.roll_number = clean_roll

        if data.department_id is not None:
            if data.department_id > 0:
                dept = db.query(Department).filter(Department.id == data.department_id).first()
                if not dept:
                    raise AppException("Selected department does not exist", status_code=400)
                user.department_id = data.department_id
                user.department = dept.name
            else:
                user.department_id = None
                user.department = None

        if data.role is not None:
            try:
                user.role = UserRole(data.role.upper())
            except ValueError:
                raise AppException(f"Invalid user role '{data.role}'", status_code=400)

        if data.password is not None and data.password.strip():
            from app.config.security import hash_password
            user.password_hash = hash_password(data.password.strip())

        if data.is_active is not None:
            user.is_active = data.is_active
            if data.is_active is False:
                from app.db.models.refresh_token import RefreshToken
                db.query(RefreshToken).filter(RefreshToken.user_id == user.id).update({"revoked": True})

        db.commit()
        db.refresh(user)

        res = AdminService.get_users(db, page=1, limit=1, search=user.email or str(user.id))
        if res.users:
            for item in res.users:
                if item.id == user.id:
                    return item
        dept = db.query(Department).filter(Department.id == user.department_id).first() if user.department_id else None
        return AdminUserItem(
            id=user.id,
            email=user.email,
            phone=getattr(user, "phone", None),
            full_name=user.full_name,
            roll_number=user.roll_number,
            department_id=user.department_id,
            department_name=dept.name if dept else user.department,
            role=user.role.value if hasattr(user.role, "value") else str(user.role),
            is_active=user.is_active,
            total_orders=0,
            total_pages=0,
            total_spent=0.0,
            created_at=user.created_at.strftime("%Y-%m-%d %H:%M:%S") if user.created_at else "",
        )

    @staticmethod
    def create_user(data: AdminUserCreate, db: Session) -> AdminUserItem:
        from app.db.models.user import UserRole
        from app.config.security import hash_password
        import secrets

        clean_email = data.email.strip().lower()
        if not clean_email or "@" not in clean_email:
            raise AppException("A valid email address is required.", status_code=400)

        existing = db.query(User).filter(User.email == clean_email).first()
        if existing:
            raise AppException("A user with this email address already exists.", status_code=400)

        clean_roll = data.roll_number.strip() if data.roll_number else None
        if clean_roll:
            existing_roll = db.query(User).filter(User.roll_number == clean_roll).first()
            if existing_roll:
                raise AppException(f"A user with roll number '{clean_roll}' already exists.", status_code=400)

        dept_id = None
        dept_name = None
        if data.department_id and data.department_id > 0:
            dept = db.query(Department).filter(Department.id == data.department_id).first()
            if not dept:
                raise AppException("Selected department does not exist.", status_code=400)
            dept_id = dept.id
            dept_name = dept.name

        try:
            role_val = UserRole(data.role.upper()) if data.role else UserRole.USER
        except ValueError:
            raise AppException(f"Invalid user role '{data.role}'", status_code=400)

        raw_password = data.password.strip() if data.password else secrets.token_urlsafe(16)
        pw_hash = hash_password(raw_password)

        new_user = User(
            email=clean_email,
            full_name=data.full_name.strip(),
            roll_number=clean_roll,
            department_id=dept_id,
            department=dept_name,
            password_hash=pw_hash,
            role=role_val,
            is_active=data.is_active,
        )
        db.add(new_user)
        db.commit()
        db.refresh(new_user)

        return AdminUserItem(
            id=new_user.id,
            email=new_user.email,
            phone=getattr(new_user, "phone", None),
            full_name=new_user.full_name,
            roll_number=new_user.roll_number,
            department_id=new_user.department_id,
            department_name=dept_name,
            role=new_user.role.value if hasattr(new_user.role, "value") else str(new_user.role),
            is_active=new_user.is_active,
            total_orders=0,
            total_pages=0,
            total_spent=0.0,
            created_at=new_user.created_at.strftime("%Y-%m-%d %H:%M:%S") if new_user.created_at else "",
        )

    @staticmethod
    def delete_user(user_id: int, current_admin_id: int, db: Session) -> bool:
        if user_id == current_admin_id:
            raise AppException("Cannot delete your own administrator account.", status_code=400)

        user = db.query(User).filter(User.id == user_id).first()
        if not user:
            raise AppException(f"User with ID {user_id} not found.", status_code=404)

        from app.db.models.idempotency import IdempotencyKey
        db.query(IdempotencyKey).filter(IdempotencyKey.user_id == user_id).delete(synchronize_session=False)

        db.delete(user)
        db.commit()
        return True

    @staticmethod
    def get_settings(current_admin_email: str) -> AdminSystemSettings:
        from app.config.settings import settings
        return AdminSystemSettings(
            app_name=settings.APP_NAME,
            per_page_rate=settings.PER_PAGE_RATE,
            base_fee=settings.BASE_FEE,
            max_upload_mb=settings.MAX_UPLOAD_MB,
            otp_ttl_minutes=settings.OTP_TTL_MINUTES,
            max_otp_attempts=settings.MAX_OTP_ATTEMPTS,
            max_print_retries=settings.MAX_PRINT_RETRIES,
            heartbeat_timeout_seconds=settings.HEARTBEAT_TIMEOUT_SECONDS,
            email_provider=settings.EMAIL_PROVIDER,
            smtp_host=settings.SMTP_HOST,
            smtp_user=settings.SMTP_USER,
            current_admin_email=current_admin_email,
        )

    @staticmethod
    def update_settings(data: AdminSystemSettingsUpdate, current_admin_email: str) -> AdminSystemSettings:
        from app.config.settings import settings
        if data.per_page_rate is not None:
            if data.per_page_rate < 0:
                raise AppException("Per page rate must be non-negative", status_code=400)
            settings.PER_PAGE_RATE = float(data.per_page_rate)

        if data.base_fee is not None:
            if data.base_fee < 0:
                raise AppException("Base fee must be non-negative", status_code=400)
            settings.BASE_FEE = float(data.base_fee)

        if data.max_upload_mb is not None:
            if data.max_upload_mb <= 0:
                raise AppException("Max upload MB must be greater than zero", status_code=400)
            settings.MAX_UPLOAD_MB = int(data.max_upload_mb)

        if data.otp_ttl_minutes is not None:
            if data.otp_ttl_minutes <= 0:
                raise AppException("OTP TTL must be greater than zero", status_code=400)
            settings.OTP_TTL_MINUTES = int(data.otp_ttl_minutes)

        if data.max_otp_attempts is not None:
            if data.max_otp_attempts <= 0:
                raise AppException("Max OTP attempts must be greater than zero", status_code=400)
            settings.MAX_OTP_ATTEMPTS = int(data.max_otp_attempts)

        if data.max_print_retries is not None:
            if data.max_print_retries < 0:
                raise AppException("Max print retries cannot be negative", status_code=400)
            settings.MAX_PRINT_RETRIES = int(data.max_print_retries)

        if data.heartbeat_timeout_seconds is not None:
            if data.heartbeat_timeout_seconds < 5:
                raise AppException("Heartbeat timeout must be at least 5 seconds", status_code=400)
            settings.HEARTBEAT_TIMEOUT_SECONDS = int(data.heartbeat_timeout_seconds)

        return AdminService.get_settings(current_admin_email)

    @staticmethod
    def get_database_status(db: Session) -> AdminDatabaseStatusResponse:
        from app.config.settings import settings
        start = time.perf_counter()
        db.execute(text("SELECT 1"))
        latency_ms = round((time.perf_counter() - start) * 1000, 2)

        # Mask password in URL
        raw_url = settings.DATABASE_URL
        if "@" in raw_url:
            masked_url = raw_url.split("://")[0] + "://***:***@" + raw_url.split("@")[-1]
        else:
            masked_url = raw_url

        return AdminDatabaseStatusResponse(
            status="CONNECTED",
            database_url_masked=masked_url,
            environment=settings.ENVIRONMENT,
            latency_ms=latency_ms,
            total_users=db.query(User).count(),
            total_departments=db.query(Department).count(),
            total_printers=db.query(Printer).count(),
            total_orders=db.query(Order).count(),
            total_print_jobs=db.query(PrintJob).count(),
        )

    @staticmethod
    def generate_roster_template() -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Student Roster Template"
        ws.views.sheetView[0].showGridLines = True

        header_fill = PatternFill(start_color="1A56DB", end_color="1A56DB", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=11)
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        headers = ["Roll Number", "Full Name", "Email", "Phone", "Department Code", "Department Name", "Role"]
        ws.append(headers)
        for col_idx in range(1, 8):
            c = ws.cell(row=1, column=col_idx)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center")

        sample_rows = [
            ["2026-CS-001", "Jane Doe", "jane.doe@campus.edu", "9876543210", "CS", "Computer Science", "USER"],
            ["2026-MECH-002", "John Smith", "john.smith@campus.edu", "9876543211", "MECH", "Mechanical Engineering", "USER"],
            ["FAC-EE-101", "Dr. Alan Turing", "alan.turing@campus.edu", "9876543212", "EE", "Electrical Engineering", "USER"],
        ]

        for r_idx, row in enumerate(sample_rows, start=2):
            ws.append(row)
            for col_idx in range(1, 8):
                c = ws.cell(row=r_idx, column=col_idx)
                c.font = regular_font
                c.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @staticmethod
    def import_roster(file_bytes: bytes, filename: str, db: Session) -> UserImportSummary:
        records = []
        is_xlsx = filename.lower().endswith(".xlsx") or filename.lower().endswith(".xls")

        if is_xlsx:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                return UserImportSummary(total_rows=0, created=0, updated=0, failed=0, errors=["The uploaded spreadsheet is empty."])
            header = [str(h or "").strip().lower() for h in rows[0]]
            for row in rows[1:]:
                if not any(row):
                    continue
                row_dict = {}
                for idx, col_val in enumerate(row):
                    if idx < len(header):
                        row_dict[header[idx]] = str(col_val or "").strip()
                records.append(row_dict)
        else:
            text_content = file_bytes.decode("utf-8-sig", errors="replace")
            reader = csv.DictReader(io.StringIO(text_content))
            for r in reader:
                records.append({k.strip().lower(): str(v or "").strip() for k, v in r.items() if k})

        if not records:
            return UserImportSummary(total_rows=0, created=0, updated=0, failed=0, errors=["No valid records found in file."])

        created_cnt = 0
        updated_cnt = 0
        failed_cnt = 0
        errors: List[str] = []

        # Cache existing departments
        depts_by_code = {d.code.upper(): d for d in db.query(Department).all()}
        depts_by_name = {d.name.lower(): d for d in db.query(Department).all()}

        for row_idx, r in enumerate(records, start=2):
            email = r.get("email", "")
            full_name = r.get("full name", r.get("name", ""))
            roll_number = r.get("roll number", r.get("roll", r.get("roll_number", "")))
            phone = r.get("phone", r.get("mobile", ""))
            dept_code = r.get("department code", r.get("dept_code", r.get("dept", ""))).upper()
            dept_name = r.get("department name", r.get("department", ""))
            role_str = r.get("role", "USER").upper()

            if not email or "@" not in email:
                failed_cnt += 1
                errors.append(f"Row {row_idx}: Missing or invalid email '{email}'")
                continue

            target_dept = None
            if dept_code and dept_code in depts_by_code:
                target_dept = depts_by_code[dept_code]
            elif dept_name and dept_name.lower() in depts_by_name:
                target_dept = depts_by_name[dept_name.lower()]
            elif dept_code or dept_name:
                # Auto-create missing department
                new_code = dept_code or dept_name[:4].upper()
                new_name = dept_name or dept_code
                target_dept = Department(code=new_code, name=new_name, description=f"Imported Department {new_name}")
                db.add(target_dept)
                db.flush()
                depts_by_code[new_code] = target_dept
                depts_by_name[new_name.lower()] = target_dept

            # Check if user exists by email or roll_number
            existing_user = db.query(User).filter(User.email == email).first()
            if not existing_user and roll_number:
                existing_user = db.query(User).filter(User.roll_number == roll_number).first()

            role_val = UserRole.ADMIN if role_str == "ADMIN" else UserRole.USER

            if existing_user:
                if full_name:
                    existing_user.full_name = full_name
                if roll_number:
                    existing_user.roll_number = roll_number
                if target_dept:
                    existing_user.department_id = target_dept.id
                    existing_user.department = target_dept.name
                existing_user.role = role_val
                updated_cnt += 1
            else:
                default_pw = secrets.token_urlsafe(16)
                new_user = User(
                    email=email,
                    full_name=full_name or email.split("@")[0],
                    roll_number=roll_number or None,
                    password_hash=hash_password(default_pw),
                    department_id=target_dept.id if target_dept else None,
                    department=target_dept.name if target_dept else None,
                    role=role_val,
                    is_active=True,
                )
                db.add(new_user)
                created_cnt += 1

        db.commit()

        return UserImportSummary(
            total_rows=len(records),
            created=created_cnt,
            updated=updated_cnt,
            failed=failed_cnt,
            errors=errors[:15],
        )

admin_service = AdminService()

